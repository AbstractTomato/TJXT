from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import build_points_workflow_doc as base


ROOT = Path(__file__).resolve().parent
ASSET_DIR = ROOT / "assets"
OUT_DOCX = ROOT / "天机学堂已开发业务流程总览.docx"


def add_title(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("天机学堂已开发业务流程总览")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = base.DARK_BLUE
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("覆盖范围：课表、学习记录、学习计划、互动问答、点赞、签到与积分")
    run.font.size = Pt(10)
    run.font.color.rgb = base.GRAY
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_note(doc, text):
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    base.set_cell_shading(cell, "FFF4E5")
    base.set_cell_border(cell, "E0B36A")
    base.set_cell_text(cell, text, size=10)
    doc.add_paragraph()


def create_diagrams():
    base.draw_flowchart(
        ASSET_DIR / "all_overview_flow.png",
        "项目业务总览：用户行为、MQ、Redis 与数据库",
        {
            "trade": {"rect": (70, 125, 270, 110), "text": "订单支付/退款\nORDER MQ", "fill": "#EAF2FF"},
            "lesson": {"rect": (440, 125, 280, 110), "text": "课表服务\nlearning_lesson", "fill": "#EAF7EA"},
            "record": {"rect": (820, 125, 280, 110), "text": "学习记录\nlearning_record", "fill": "#EAF7EA"},
            "qa": {"rect": (70, 340, 270, 120), "text": "互动问答\n问题/回答/评论", "fill": "#FFF4E5"},
            "remark": {"rect": (440, 340, 280, 120), "text": "点赞服务\nRedis Set/ZSet", "fill": "#EEF2FF"},
            "points": {"rect": (820, 340, 280, 120), "text": "积分服务\npoints_record", "fill": "#EAF7EA"},
            "query": {"rect": (1190, 230, 240, 125), "text": "前端查询\n课表/问答/积分", "fill": "#F5F7FA"},
        },
        [
            ("trade", "lesson", None),
            ("lesson", "record", None),
            ("record", "points", None),
            ("qa", "remark", None),
            ("remark", "qa", None),
            ("qa", "points", None),
            ("lesson", "query", None),
            ("qa", "query", None),
            ("points", "query", None),
        ],
        height=700,
    )

    base.draw_flowchart(
        ASSET_DIR / "lesson_flow.png",
        "课表业务：订单驱动课表，用户查询/删除课表",
        {
            "pay": {"rect": (70, 150, 250, 100), "text": "支付成功 MQ\nORDER_PAY", "fill": "#EAF2FF"},
            "listener": {"rect": (410, 150, 280, 100), "text": "LessonChangeListener\n幂等去重", "fill": "#F5F7FA"},
            "add": {"rect": (780, 150, 280, 100), "text": "addUserLessons\n批量创建课表", "fill": "#EAF7EA"},
            "query": {"rect": (1160, 150, 260, 100), "text": "查询课表\n/page /{courseId}", "fill": "#EAF7EA"},
            "refund": {"rect": (70, 420, 250, 100), "text": "退款成功 MQ\nORDER_REFUND", "fill": "#FFF4E5"},
            "delete": {"rect": (410, 420, 280, 100), "text": "deleteCourseFromLesson\n删除课表课程", "fill": "#FFF4E5"},
            "valid": {"rect": (780, 420, 280, 100), "text": "valid/count\n远程校验与统计", "fill": "#EEF2FF"},
        },
        [
            ("pay", "listener", None),
            ("listener", "add", None),
            ("add", "query", None),
            ("refund", "delete", None),
            ("query", "valid", None),
        ],
        height=700,
    )

    base.draw_flowchart(
        ASSET_DIR / "learning_record_flow.png",
        "学习记录：视频/考试提交、课表进度与学习积分",
        {
            "api": {"rect": (70, 170, 260, 100), "text": "POST /learning-records\n提交学习记录", "fill": "#EAF2FF"},
            "type": {"rect": (420, 170, 260, 100), "text": "判断小节类型\nVIDEO / EXAM", "fill": "#FFF4E5"},
            "video": {"rect": (770, 100, 270, 100), "text": "视频处理\n进度未完成则延迟写", "fill": "#F5F7FA"},
            "exam": {"rect": (770, 270, 270, 100), "text": "考试处理\n提交即完成", "fill": "#F5F7FA"},
            "lesson": {"rect": (1130, 170, 280, 100), "text": "首次完成\n更新课表进度", "fill": "#EAF7EA"},
            "mq": {"rect": (1130, 420, 280, 100), "text": "发送 LEARN_SECTION\n积分 +10", "fill": "#EEF2FF"},
        },
        [
            ("api", "type", None),
            ("type", "video", None),
            ("type", "exam", None),
            ("video", "lesson", None),
            ("exam", "lesson", None),
            ("lesson", "mq", None),
        ],
        height=700,
    )

    base.draw_flowchart(
        ASSET_DIR / "qa_flow.png",
        "互动问答：问题、回答、评论、管理端审核",
        {
            "ask": {"rect": (70, 130, 250, 100), "text": "用户提问\nPOST /questions", "fill": "#EAF2FF"},
            "page": {"rect": (410, 130, 250, 100), "text": "用户端查询\n/page /{id}", "fill": "#EAF7EA"},
            "reply": {"rect": (750, 130, 280, 100), "text": "回答/评论\nPOST /replies", "fill": "#FFF4E5"},
            "points": {"rect": (1120, 130, 280, 100), "text": "学生回答\n发送 WRITE_REPLY", "fill": "#EEF2FF"},
            "admin": {"rect": (250, 410, 280, 100), "text": "管理端分页/详情\n/admin/questions", "fill": "#F5F7FA"},
            "hide": {"rect": (630, 410, 300, 100), "text": "隐藏问题/回复\n同步 hidden 状态", "fill": "#FFF4E5"},
            "status": {"rect": (1030, 410, 300, 100), "text": "查看详情后\n状态改为 CHECKED", "fill": "#EAF7EA"},
        },
        [
            ("ask", "page", None),
            ("page", "reply", None),
            ("reply", "points", None),
            ("admin", "hide", None),
            ("admin", "status", None),
        ],
        height=700,
    )

    base.draw_flowchart(
        ASSET_DIR / "remark_flow.png",
        "点赞业务：高频状态写 Redis，批量同步点赞数",
        {
            "api": {"rect": (70, 170, 260, 100), "text": "POST /likes\n点赞/取消点赞", "fill": "#EAF2FF"},
            "set": {"rect": (420, 170, 270, 100), "text": "Redis Set\n保存用户点赞状态", "fill": "#EAF7EA"},
            "zset": {"rect": (780, 170, 270, 100), "text": "Redis ZSet\n记录业务点赞数变化", "fill": "#EEF2FF"},
            "task": {"rect": (1130, 170, 270, 100), "text": "定时任务\npopMin 批量读取", "fill": "#FFF4E5"},
            "mq": {"rect": (780, 420, 270, 100), "text": "发送 LIKE MQ\nLikedTimesDTO 列表", "fill": "#F5F7FA"},
            "learning": {"rect": (420, 420, 270, 100), "text": "Learning 监听\n更新 liked_times", "fill": "#EAF7EA"},
        },
        [
            ("api", "set", None),
            ("set", "zset", None),
            ("zset", "task", None),
            ("task", "mq", None),
            ("mq", "learning", None),
        ],
        height=700,
    )


def add_endpoint_table(doc):
    base.add_table(
        doc,
        ["业务域", "接口/触发点", "状态", "说明"],
        [
            ["课表", "GET /lessons/page", "已实现", "查询当前用户课表分页，并补充课程名称、封面、小节数。"],
            ["课表", "GET /lessons/{courseId}", "已实现", "查询当前用户指定课程学习状态。"],
            ["课表", "DELETE /lessons/{courseId}", "已实现", "用户主动删除课表课程。"],
            ["课表", "GET /lessons/{courseId}/valid", "已实现", "Feign 校验课程是否是有效课表课程，过期返回 null。"],
            ["课表", "GET /lessons/{courseId}/count", "已实现", "统计课程学习人数。"],
            ["学习计划", "POST /lessons/plans", "已实现", "创建学习计划，更新 weekFreq 与 planStatus。"],
            ["学习计划", "POST /lessons/plans", "需调整", "查询学习计划当前也使用 POST /plans，和创建计划路径冲突，后续建议改成 GET /plans。"],
            ["学习记录", "GET /learning-records/course/{courseId}", "已实现", "查询课程下课表与学习记录。"],
            ["学习记录", "POST /learning-records", "已实现", "提交视频/考试学习记录，首次完成后更新课表并发学习积分 MQ。"],
            ["问答", "POST /questions", "已实现", "用户新增问题。"],
            ["问答", "GET /questions/page", "已实现", "用户端分页查询问题，支持只看自己、课程/小节筛选。"],
            ["问答", "GET /questions/{id}", "已实现", "用户端查询问题详情，隐藏问题不可见。"],
            ["回复", "POST /replies", "已实现", "新增回答或评论；学生回答问题后发送问答积分 MQ。"],
            ["回复", "GET /replies/page", "已实现", "分页查询回答或评论，隐藏数据用户端不可见。"],
            ["管理端问答", "GET /admin/questions/page", "已实现", "按课程名、状态、时间分页查询问题。"],
            ["管理端问答", "PUT /admin/questions/{id}/hidden/{hidden}", "已实现", "隐藏/显示问题，并同步隐藏回复评论。"],
            ["管理端问答", "GET /admin/questions/{id}", "已实现", "查询问题详情，并将状态更新为已查看。"],
            ["管理端回复", "GET /admin/replies/page", "已实现", "管理端分页查询回答/评论。"],
            ["管理端回复", "PUT /admin/replies/{id}/hidden/{hidden}", "已实现", "隐藏/显示回答；回答隐藏时同步处理其评论。"],
            ["点赞", "POST /likes", "已实现", "点赞/取消点赞，状态写 Redis Set。"],
            ["点赞", "GET /likes/list", "已实现", "批量查询当前用户对业务 id 的点赞状态。"],
            ["签到", "POST /sign-records", "已实现", "签到、连续签到奖励、发送签到积分 MQ。"],
            ["签到", "GET /sign-records", "已实现", "查询本月签到结果。"],
            ["积分", "GET /points/today", "已实现", "查询当前用户今日积分统计。"],
            ["排行榜", "GET /boards 等", "暂缓", "当前未展开，后续再接 Redis ZSet 和赛季历史表。"],
        ],
        [2.0, 4.0, 1.6, 8.0],
    )


def build_doc():
    create_diagrams()
    doc = Document()
    base.set_doc_styles(doc)
    add_title(doc)

    doc.add_heading("1. 文档范围", level=1)
    doc.add_paragraph(
        "本文档整理截至当前已经开发和讨论过的核心业务流程。重点覆盖 tj-learning 与 tj-remark 两个模块，"
        "包括课表、学习记录、学习计划、互动问答、点赞、签到与积分。排行榜/赛季榜单已按当前计划暂缓，仅保留后续提醒。"
    )
    add_note(
        doc,
        "阅读提示：本文档是业务流程复习稿，不替代接口文档。表格用于快速定位接口，流程图用于理解数据如何在 Controller、Service、MQ、Redis、MySQL 之间流转。",
    )

    doc.add_heading("2. 接口与业务状态总表", level=1)
    add_endpoint_table(doc)

    doc.add_heading("3. 项目总流程图", level=1)
    doc.add_paragraph(
        "当前项目的核心特点是：订单通过 MQ 驱动课表变化；学习、签到、回答等行为通过 MQ 驱动积分；"
        "点赞高频状态先进入 Redis，再批量同步回学习服务。"
    )
    doc.add_picture(str(ASSET_DIR / "all_overview_flow.png"), width=Inches(6.5))

    doc.add_heading("4. 课表业务流程", level=1)
    doc.add_picture(str(ASSET_DIR / "lesson_flow.png"), width=Inches(6.5))
    base.add_table(
        doc,
        ["流程", "关键类", "核心逻辑"],
        [
            ["支付成功添加课表", "LessonChangeListener + LearningLessonServiceImpl.addUserLessons", "监听订单支付 MQ；校验消息；对 courseIds 去重；查询用户课表已有课程做幂等过滤；批量保存 LearningLesson。"],
            ["退款删除课表", "LessonChangeListener.listenCourseRefund", "监听退款 MQ；校验订单数据；调用 deleteCourseFromLesson 删除对应课程。删除天然幂等。"],
            ["用户查询课表", "queryMyLessons", "根据 UserContext 查询当前用户课表分页，再调用 CourseClient 批量补充课程信息。"],
            ["课表有效性校验", "isLessonValid", "按 userId + courseId 查课表；若不存在或已过期，返回 null；否则返回 lessonId。"],
            ["学习人数统计", "countLearningPersonByCourse", "统计指定课程下状态为未开始、学习中、已学完的课表数量。"],
        ],
        [3.0, 4.4, 7.2],
    )

    doc.add_heading("5. 学习记录与学习进度流程", level=1)
    doc.add_picture(str(ASSET_DIR / "learning_record_flow.png"), width=Inches(6.5))
    base.add_table(
        doc,
        ["环节", "视频小节", "考试小节"],
        [
            ["是否存在旧记录", "先查延迟缓存，再查 learning_record。", "直接新增考试记录。"],
            ["完成条件", "旧记录未完成，且本次 moment 达到 duration 的一半。", "提交考试记录即视为完成。"],
            ["未完成处理", "写入延迟任务，后续异步更新播放进度。", "无。"],
            ["首次完成处理", "更新 finished、finishTime，清理缓存。", "保存 finished=true、finishTime=commitTime。"],
            ["后续动作", "更新课表学习进度，发送 LEARN_SECTION 积分消息。", "更新课表学习进度，发送 LEARN_SECTION 积分消息。"],
        ],
        [2.2, 6.2, 6.2],
    )

    doc.add_heading("6. 学习计划流程", level=1)
    doc.add_paragraph(
        "学习计划本质上是更新 learning_lesson 表中的 weekFreq 与 planStatus。查询学习计划时，会统计本周已完成小节数、"
        "本周计划小节数，并分页返回正在计划中且未完成的课程。"
    )
    base.add_table(
        doc,
        ["流程", "关键逻辑"],
        [
            ["创建学习计划", "根据当前用户与 courseId 查询课表；设置 weekFreq；将 planStatus 设置为 PLAN_RUNNING。"],
            ["查询本周计划汇总", "统计本周 finished=true 且 finishTime 在本周范围内的学习记录数量。"],
            ["查询本周计划总量", "调用 LearningLessonMapper.queryTotalPlan(userId) 汇总计划频次。"],
            ["分页返回计划课程", "筛选当前用户、PLAN_RUNNING、课程状态未开始或学习中；补充课程名、小节数和本周已学小节数。"],
            ["当前注意点", "Controller 中创建计划和查询计划都标注为 POST /lessons/plans，运行时可能路径冲突；后续建议查询改为 GET /lessons/plans。"],
        ],
        [3.4, 11.0],
    )

    doc.add_heading("7. 互动问答流程", level=1)
    doc.add_picture(str(ASSET_DIR / "qa_flow.png"), width=Inches(6.5))
    base.add_table(
        doc,
        ["业务", "关键规则"],
        [
            ["新增问题", "当前用户从 UserContext 获取 userId，QuestionFormDTO 拷贝为 InteractionQuestion 后保存。"],
            ["用户端问题分页", "courseId 与 sectionId 不能同时为空；默认过滤 hidden=false；匿名问题不查用户信息。"],
            ["问题详情", "问题不存在或 hidden=true 时返回 null；非匿名时补充提问人信息。"],
            ["新增回答/评论", "answerId 为空表示回答；answerId 非空表示评论。回答更新 latestAnswerId 与 answerTimes；评论更新对应回答 replyTimes。"],
            ["问答积分", "当前版本：学生回答问题才发送 WRITE_REPLY，评论不加分。若要完全对齐参考仓库，可改成学生回答或评论都发 MQ。"],
            ["管理端隐藏问题", "更新问题 hidden，同时按 questionId 同步更新该问题下回复/评论 hidden。"],
            ["管理端隐藏回复", "隐藏回答时同步隐藏其评论；隐藏评论时只处理当前评论。"],
            ["管理端详情", "补充用户、课程、分类、章节、老师信息，并将问题状态更新为 CHECKED。"],
        ],
        [3.0, 11.4],
    )

    doc.add_heading("8. 点赞业务流程", level=1)
    doc.add_picture(str(ASSET_DIR / "remark_flow.png"), width=Inches(6.5))
    base.add_table(
        doc,
        ["环节", "说明"],
        [
            ["点赞/取消点赞", "前端传 bizId、bizType、liked；liked=true 时 SADD，liked=false 时 SREM。"],
            ["点赞状态", "Redis Set 保存某业务下所有点赞用户，key 类似 likes:biz:{bizId}。"],
            ["点赞数变更", "执行成功后统计 Set size，将业务 id 与点赞总数写入 Redis ZSet。"],
            ["批量同步", "LikedTimesCheckTask 每 20 秒扫描 QA、NOTE 类型，popMin 批量取出变化数据。"],
            ["发送 MQ", "封装 LikedTimesDTO 列表，根据业务类型生成 routing key，发送到点赞交换机。"],
            ["学习服务落库", "LikedTimesChangeListener 监听 QA 点赞数变化，批量更新 interaction_reply.liked_times。"],
            ["查询点赞状态", "GET /likes/list 用 pipeline 批量判断当前用户是否在各业务 Set 中。"],
        ],
        [3.0, 11.4],
    )

    doc.add_heading("9. 签到与积分流程", level=1)
    doc.add_paragraph("这部分沿用上一份积分文档的核心结论，放在总文档中便于统一复习。")
    doc.add_picture(str(ASSET_DIR / "overall_flow.png"), width=Inches(6.5))
    doc.add_picture(str(ASSET_DIR / "add_points_flow.png"), width=Inches(6.5))
    doc.add_picture(str(ASSET_DIR / "today_query_flow.png"), width=Inches(6.5))
    base.add_table(
        doc,
        ["业务", "关键逻辑"],
        [
            ["签到", "Redis BitMap 记录本月签到；连续 7/14/28 天分别奖励 10/20/40 分；发送 SIGN_IN。"],
            ["学习积分", "首次学完小节后发送 LEARN_SECTION，积分 +10，每日上限由 LEARNING 枚举控制。"],
            ["问答积分", "学生回答问题后发送 WRITE_REPLY，积分 +5，每日上限由 QA 枚举控制。"],
            ["积分入账", "LearningPointsListener 根据 routing key 识别积分类型；PointsRecordServiceImpl 判断每日上限后写 points_record。"],
            ["今日积分查询", "GET /points/today 按当前用户、今日时间范围、type 分组统计 points。"],
        ],
        [3.0, 11.4],
    )

    doc.add_heading("10. MQ 与异步链路汇总", level=1)
    base.add_table(
        doc,
        ["来源", "Exchange", "Routing Key", "消费者", "结果"],
        [
            ["订单支付", "ORDER_EXCHANGE", "ORDER_PAY_KEY", "LessonChangeListener", "幂等添加课表。"],
            ["订单退款", "ORDER_EXCHANGE", "ORDER_REFUND_KEY", "LessonChangeListener", "删除课表课程。"],
            ["签到成功", "LEARNING_EXCHANGE", "SIGN_IN", "LearningPointsListener", "写 SIGN 积分记录。"],
            ["学完小节", "LEARNING_EXCHANGE", "LEARN_SECTION", "LearningPointsListener", "写 LEARNING 积分记录。"],
            ["学生回答", "LEARNING_EXCHANGE", "WRITE_REPLY", "LearningPointsListener", "写 QA 积分记录。"],
            ["点赞数变化", "LIKE_RECORD_EXCHANGE", "QA_LIKED_TIMES_KEY", "LikedTimesChangeListener", "更新回答点赞数。"],
        ],
        [2.4, 3.0, 3.2, 3.2, 4.6],
    )

    doc.add_heading("11. 当前遗留与后续提醒", level=1)
    base.add_bullet(doc, "排行榜与赛季榜单暂缓：PointsBoardController、PointsBoardSeasonController 仍是生成器空壳。")
    base.add_bullet(doc, "学习计划查询接口建议调整：当前创建计划与查询计划都是 POST /lessons/plans，后续应改成不同 HTTP 方法或不同路径。")
    base.add_bullet(doc, "积分入账暂未接 Redis ZSet 排行榜累加；后续写排行榜时再补 POINTS_BOARD_KEY_PREFIX 与 incrementScore。")
    base.add_bullet(doc, "点赞同步目前使用定时任务，每 20 秒批量发送 MQ；如果后续要改延迟任务，需要重新设计触发和去重策略。")
    base.add_bullet(doc, "部分生成代码注释在终端显示乱码，建议后续统一文件编码为 UTF-8，方便长期维护。")

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("生成说明：本文档基于当前本地代码与此前开发过程整理。")
    run.font.size = Pt(9)
    run.font.color.rgb = base.GRAY
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    doc.save(OUT_DOCX)
    return OUT_DOCX


if __name__ == "__main__":
    print(build_doc())
