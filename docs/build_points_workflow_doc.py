from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent
ASSET_DIR = ROOT / "assets"
OUT_DOCX = ROOT / "积分业务流程说明.docx"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
GRAY = RGBColor(85, 85, 85)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
WHITE = "FFFFFF"
BORDER = "AEB7C2"


def font(size, bold=False):
    candidates = [
        "C:/Windows/Fonts/msyh.ttc",
        "C:/Windows/Fonts/simhei.ttf",
        "C:/Windows/Fonts/simsun.ttc",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size=size)
    return ImageFont.load_default()


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, color=None, size=10):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_table(doc, headers, rows, widths_cm):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    for i, width in enumerate(widths_cm):
        table.columns[i].width = Cm(width)
    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        set_cell_text(header_cells[i], header, bold=True, color=DARK_BLUE)
        set_cell_shading(header_cells[i], LIGHT_BLUE)
        set_cell_border(header_cells[i])
        header_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value), size=9)
            set_cell_border(cells[i])
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    doc.add_paragraph()
    return table


def draw_flowchart(path, title, nodes, edges, width=1500, height=760):
    img = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(img)
    title_font = font(38, bold=True)
    node_font = font(24)
    small_font = font(20)

    d.text((50, 32), title, fill="#1F4D78", font=title_font)

    def center(rect):
        x, y, w, h = rect
        return x + w / 2, y + h / 2

    def draw_arrow(a, b, label=None):
        ax, ay = center(nodes[a]["rect"])
        bx, by = center(nodes[b]["rect"])
        ar = nodes[a]["rect"]
        br = nodes[b]["rect"]
        if abs(bx - ax) >= abs(by - ay):
            start = (ar[0] + ar[2] if bx > ax else ar[0], ay)
            end = (br[0] if bx > ax else br[0] + br[2], by)
        else:
            start = (ax, ar[1] + ar[3] if by > ay else ar[1])
            end = (bx, br[1] if by > ay else br[1] + br[3])
        d.line([start, end], fill="#546A7B", width=4)
        ex, ey = end
        sx, sy = start
        if abs(ex - sx) >= abs(ey - sy):
            sign = 1 if ex > sx else -1
            arrow = [(ex, ey), (ex - sign * 18, ey - 10), (ex - sign * 18, ey + 10)]
        else:
            sign = 1 if ey > sy else -1
            arrow = [(ex, ey), (ex - 10, ey - sign * 18), (ex + 10, ey - sign * 18)]
        d.polygon(arrow, fill="#546A7B")
        if label:
            lx = (start[0] + end[0]) / 2
            ly = (start[1] + end[1]) / 2 - 28
            d.rounded_rectangle((lx - 85, ly - 18, lx + 85, ly + 18), radius=8, fill="#FFFFFF", outline="#D0D7DE")
            d.text((lx - 72, ly - 13), label, fill="#4A5568", font=small_font)

    for a, b, label in edges:
        draw_arrow(a, b, label)

    for key, item in nodes.items():
        x, y, w, h = item["rect"]
        fill = item.get("fill", "#E8EEF5")
        d.rounded_rectangle((x, y, x + w, y + h), radius=20, fill=fill, outline="#8AA4C2", width=3)
        lines = item["text"].split("\n")
        total_h = len(lines) * 30
        yy = y + (h - total_h) / 2
        for line in lines:
            bbox = d.textbbox((0, 0), line, font=node_font)
            d.text((x + (w - (bbox[2] - bbox[0])) / 2, yy), line, fill="#0B2545", font=node_font)
            yy += 30
    img.save(path)


def create_diagrams():
    draw_flowchart(
        ASSET_DIR / "overall_flow.png",
        "积分业务总览：行为 -> MQ -> 积分明细",
        {
            "sign": {"rect": (70, 135, 290, 120), "text": "每日签到\nPOST /sign-records\nSIGN_IN", "fill": "#EAF7EA"},
            "learn": {"rect": (70, 315, 290, 120), "text": "学完小节\n/learning-records\nLEARN_SECTION", "fill": "#EAF2FF"},
            "reply": {"rect": (70, 495, 290, 120), "text": "学生回答\nPOST /replies\nWRITE_REPLY", "fill": "#FFF4E5"},
            "mq": {"rect": (470, 300, 260, 130), "text": "RabbitMQ\nlearning.exchange", "fill": "#F5F7FA"},
            "listener": {"rect": (870, 300, 280, 130), "text": "LearningPointsListener\n按 routing key 分流", "fill": "#EEF2FF"},
            "record": {"rect": (1210, 300, 230, 130), "text": "points_record\n积分明细落库", "fill": "#EAF7EA"},
        },
        [
            ("sign", "mq", None),
            ("learn", "mq", None),
            ("reply", "mq", None),
            ("mq", "listener", None),
            ("listener", "record", None),
        ],
    )

    draw_flowchart(
        ASSET_DIR / "add_points_flow.png",
        "积分入账：每日上限判断与真实积分计算",
        {
            "start": {"rect": (70, 170, 230, 90), "text": "收到积分消息\nuserId + points", "fill": "#EAF2FF"},
            "valid": {"rect": (380, 170, 230, 90), "text": "校验消息\n缺字段直接返回", "fill": "#F5F7FA"},
            "limit": {"rect": (690, 170, 250, 90), "text": "读取类型上限\nmaxPoints", "fill": "#FFF4E5"},
            "sum": {"rect": (1040, 170, 300, 90), "text": "有上限则查询\n今日同类型已得分", "fill": "#F5F7FA"},
            "calc": {"rect": (1040, 420, 300, 100), "text": "超过上限则返回\n接近上限则截断积分", "fill": "#FFF4E5"},
            "save": {"rect": (690, 420, 280, 100), "text": "保存 PointsRecord\n实际积分 realPoints", "fill": "#EAF7EA"},
            "end": {"rect": (380, 420, 170, 100), "text": "结束", "fill": "#EAF7EA"},
        },
        [
            ("start", "valid", None),
            ("valid", "limit", None),
            ("limit", "sum", None),
            ("sum", "calc", None),
            ("calc", "save", None),
            ("limit", "save", None),
            ("save", "end", None),
        ],
    )

    draw_flowchart(
        ASSET_DIR / "today_query_flow.png",
        "今日积分查询：按类型分组统计",
        {
            "api": {"rect": (90, 260, 260, 100), "text": "GET /points/today\n无请求参数", "fill": "#EAF2FF"},
            "user": {"rect": (450, 260, 230, 100), "text": "UserContext\n获取当前用户", "fill": "#F5F7FA"},
            "time": {"rect": (780, 260, 250, 100), "text": "计算今日区间\n00:00:00 ~ 23:59:59", "fill": "#FFF4E5"},
            "sql": {"rect": (1130, 260, 300, 100), "text": "points_record\nGROUP BY type", "fill": "#EAF7EA"},
            "vo": {"rect": (555, 500, 360, 100), "text": "PointsStatisticsVO\n类型描述 + 已得分 + 上限", "fill": "#EEF2FF"},
        },
        [
            ("api", "user", None),
            ("user", "time", None),
            ("time", "sql", None),
            ("sql", "vo", None),
        ],
        height=680,
    )


def set_doc_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25


def add_title(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("学习积分业务流程说明")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = DARK_BLUE
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("tj-learning 模块 | 当前阶段：签到、积分入账、学习积分、问答积分、今日积分查询")
    run.font.size = Pt(10)
    run.font.color.rgb = GRAY
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    run = p.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(10.5)


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(46, 46, 46)


def build_doc():
    create_diagrams()
    doc = Document()
    set_doc_styles(doc)
    add_title(doc)

    doc.add_heading("1. 当前开发结论", level=1)
    doc.add_paragraph(
        "当前阶段，排行榜与赛季榜单相关接口已按开发计划暂时搁置；非排行榜部分的积分链路已经基本打通。"
        "本文档整理已经实现或刚刚接入的业务流程，便于后续复习、联调和继续开发排行榜。"
    )
    add_table(
        doc,
        ["模块", "接口/触发点", "当前状态", "说明"],
        [
            ["签到", "POST /sign-records", "已实现", "Redis BitMap 记录当月签到；发送 SIGN_IN 积分消息。"],
            ["查询签到", "GET /sign-records", "已实现", "从 Redis BitMap 解析本月 1 号到今天的签到结果。"],
            ["积分入账", "MQ 消费", "已实现", "LearningPointsListener 消费消息并写入 points_record。"],
            ["学习积分", "学习记录提交后触发", "已接入", "首次学完小节后发送 LEARN_SECTION，积分值 10。"],
            ["问答积分", "学生回答后触发", "已接入", "学生回答问题后发送 WRITE_REPLY，积分值 5；评论暂不加分。"],
            ["今日积分", "GET /points/today", "已实现", "按当前用户、今日时间范围、积分类型分组统计。"],
            ["排行榜", "GET /boards 等", "暂缓", "当前不展开；后续需要 Redis ZSet 和赛季历史表。"],
        ],
        [2.0, 3.0, 1.6, 9.0],
    )

    doc.add_heading("2. 总体业务流", level=1)
    doc.add_paragraph(
        "积分不是由业务接口直接写入，而是通过 MQ 进行解耦。签到、学习、回答这些行为只负责发送积分事件；"
        "统一由 LearningPointsListener 消费消息，再调用 IPointsRecordService.addPointsRecord 完成积分入账。"
    )
    doc.add_picture(str(ASSET_DIR / "overall_flow.png"), width=Inches(6.5))

    doc.add_heading("3. 签到流程", level=1)
    doc.add_paragraph("签到包含两个接口：一个负责今日签到，一个负责查询本月签到记录。")
    add_table(
        doc,
        ["步骤", "代码位置", "关键逻辑"],
        [
            ["1", "SignRecordController.addSignRecords", "接收 POST /sign-records 请求，调用 Service。"],
            ["2", "SignRecordServiceImpl.addSignRecords", "拼接 Redis key：sign:uid:{userId}:{yyyyMM}。"],
            ["3", "Redis setBit", "offset = dayOfMonth - 1；如果旧值已经是 true，说明重复签到。"],
            ["4", "countSignDays", "用 bitField 读取当月到今天的签到位图，从低位开始统计连续签到天数。"],
            ["5", "奖励计算", "连续 7/14/28 天分别奖励 10/20/40 分，基础签到分为 1。"],
            ["6", "发送 MQ", "发送 SIGN_IN，消息体为 SignInMessage.of(userId, rewardPoints + 1)。"],
        ],
        [1.2, 4.2, 10.2],
    )
    add_code(doc, "Redis key: sign:uid:{userId}:{yyyyMM}    offset: dayOfMonth - 1")
    add_code(doc, "MQ: LEARNING_EXCHANGE + SIGN_IN + SignInMessage(userId, rewardPoints + 1)")

    doc.add_heading("4. 积分入账流程", level=1)
    doc.add_paragraph(
        "积分入账的核心点是每日上限判断。签到没有上限，因此不能把 save 写在 maxPoints > 0 的分支里。"
        "当前代码已经修正为：有上限时先判断，没有上限或未达到上限时统一保存积分明细。"
    )
    doc.add_picture(str(ASSET_DIR / "add_points_flow.png"), width=Inches(6.5))
    add_table(
        doc,
        ["积分类型", "枚举", "单次积分", "每日上限", "当前来源"],
        [
            ["课程学习", "LEARNING", "10", "50", "首次学完一个小节后发送 MQ。"],
            ["每日签到", "SIGN", "1 + 连续签到奖励", "0", "签到成功后发送 MQ；0 表示无每日上限。"],
            ["课程问答", "QA", "5", "20", "学生回答问题后发送 MQ；当前评论不加分。"],
            ["学习笔记", "NOTE", "暂未接入", "20", "后续笔记模块开发后再接。"],
            ["课程评价", "COMMENT", "暂未接入", "0", "后续评价模块开发后再接。"],
        ],
        [2.2, 2.0, 2.0, 2.0, 7.0],
    )

    doc.add_heading("5. 学习积分流程", level=1)
    doc.add_paragraph(
        "学习积分由学习记录提交接口触发。只有本次提交让小节第一次从未完成变成已完成，才会发送积分消息，避免重复刷进度导致重复加分。"
    )
    add_table(
        doc,
        ["步骤", "判断点", "说明"],
        [
            ["1", "addLearningRecord", "获取当前用户，并区分视频小节或考试小节。"],
            ["2", "handleVideoRecord / handleExamRecord", "返回 finished，表示本次是否第一次完成小节。"],
            ["3", "!finished 则 return", "不是第一次完成，不更新课表，也不给积分。"],
            ["4", "handleLearningLessonsChanges", "更新课表学习状态、已学小节数量、课程完成状态。"],
            ["5", "mqHelper.send", "发送 LEARN_SECTION，消息体为 SignInMessage.of(userId, 10)。"],
        ],
        [1.2, 4.0, 10.4],
    )
    add_code(doc, "MQ: LEARNING_EXCHANGE + LEARN_SECTION + SignInMessage(userId, 10)")

    doc.add_heading("6. 问答积分流程", level=1)
    doc.add_paragraph(
        "问答积分由新增回答接口触发。当前实现选择“学生回答问题加分，评论不加分”，因此会额外判断 answerId == null。"
        "这比参考仓库更贴近“回答问题加积分”的需求语义；如果后续要完全对齐参考仓库，可以去掉 answerId 判断。"
    )
    add_table(
        doc,
        ["步骤", "代码位置", "说明"],
        [
            ["1", "InteractionReplyServiceImpl.addReply", "保存回答或评论，并设置当前用户 id。"],
            ["2", "answerId 判断", "answerId 为空表示回答；非空表示评论。"],
            ["3", "更新问题/回答统计", "回答时更新 latestAnswerId 与 answerTimes；评论时更新 replyTimes。"],
            ["4", "学生提交状态", "学生提交后将问题状态改为 UN_CHECK。"],
            ["5", "发送积分消息", "isStudent 为 true 且 answerId 为空时发送 WRITE_REPLY，积分值 5。"],
        ],
        [1.2, 4.2, 10.2],
    )
    add_code(doc, "MQ: LEARNING_EXCHANGE + WRITE_REPLY + SignInMessage(userId, 5)")

    doc.add_heading("7. 今日积分查询流程", level=1)
    doc.add_paragraph(
        "今日积分接口用于展示当前用户当天不同积分类型的获得情况。它不需要请求参数，用户身份来自 UserContext。"
    )
    doc.add_picture(str(ASSET_DIR / "today_query_flow.png"), width=Inches(6.5))
    add_table(
        doc,
        ["字段", "含义", "来源"],
        [
            ["type", "积分类型描述，例如每日签到、课程学习、课程问答", "PointsRecordType.getDesc()"],
            ["points", "今天该类型已获得积分总和", "SUM(points)"],
            ["maxPoints", "该类型每日积分上限，0 表示无上限", "PointsRecordType.getMaxPoints()"],
        ],
        [2.0, 7.0, 5.6],
    )
    add_code(doc, "SQL 语义: SELECT type, SUM(points) AS points FROM points_record WHERE user_id = ? AND create_time BETWEEN ? AND ? GROUP BY type")

    doc.add_heading("8. 当前关键类清单", level=1)
    add_table(
        doc,
        ["类别", "类/文件", "职责"],
        [
            ["Controller", "SignRecordController", "签到与本月签到记录查询入口。"],
            ["Controller", "PointsRecordController", "今日积分查询入口。"],
            ["Service", "SignRecordServiceImpl", "Redis BitMap 签到、连续签到计算、发送签到积分 MQ。"],
            ["Service", "LearningRecordServiceImpl", "学习记录处理、首次完成小节后发送学习积分 MQ。"],
            ["Service", "InteractionReplyServiceImpl", "回答/评论保存、学生回答后发送问答积分 MQ。"],
            ["MQ", "LearningPointsListener", "监听 SIGN_IN、LEARN_SECTION、WRITE_REPLY 三类积分消息。"],
            ["Service", "PointsRecordServiceImpl", "积分上限判断、积分明细落库、今日积分统计。"],
            ["Mapper", "PointsRecordMapper", "按类型求和统计积分。"],
            ["PO/VO", "PointsRecord / PointsStatisticsVO", "积分明细实体与今日积分返回模型。"],
        ],
        [2.0, 4.2, 8.4],
    )

    doc.add_heading("9. 后续开发提醒", level=1)
    add_bullet(doc, "排行榜暂缓：后续需要给积分入账增加 Redis ZSet 累加逻辑，key 形如 boards:{yyyyMM}。")
    add_bullet(doc, "赛季榜单暂缓：后续需要 PointsBoardQuery、PointsBoardVO、PointsBoardItemVO、动态表名与历史赛季表。")
    add_bullet(doc, "当前 PointsRecordServiceImpl 中 StringRedisTemplate 暂时未使用，等排行榜接入时再使用。")
    add_bullet(doc, "如果要完全对齐参考仓库，问答积分可改为学生回答或评论都加分；当前版本只给回答加分。")
    add_bullet(doc, "联调时建议按顺序验证：签到一次、学习完成一个小节、提交一个学生回答、调用 GET /points/today。")

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("生成说明：本文档基于当前 tj-learning 代码与参考仓库业务逻辑整理。")
    run.font.size = Pt(9)
    run.font.color.rgb = GRAY
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    doc.save(OUT_DOCX)
    return OUT_DOCX


if __name__ == "__main__":
    path = build_doc()
    print(path)
